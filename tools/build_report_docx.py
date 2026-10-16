# -*- coding: utf-8 -*-
"""把评估报告的 Markdown 渲染为排版规范的 DOCX。"""
import os
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# 路径基于脚本位置推导，保证在任意工作目录下都能运行
_HERE = os.path.dirname(os.path.abspath(__file__))
_ROOT = os.path.dirname(_HERE)
_DOCS = os.path.join(_ROOT, "docs")
SRC = os.path.join(_DOCS, "代码安全与可行性评估报告.md")
OUT = os.path.join(_DOCS, "代码安全与可行性评估报告.docx")

BODY_ASCII = "Calibri"
BODY_CN = "宋体"
HEAD_ASCII = "Segoe UI"
HEAD_CN = "微软雅黑"
MONO = "Consolas"

PAGE_W = 21.0  # A4 宽 cm
MARGIN = 2.2
CONTENT_W = PAGE_W - 2 * MARGIN

INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]*\))")
LINK_RE = re.compile(r"^\[([^\]]+)\]\(([^)]*)\)$")


def shade(element, fill):
    rPr = element.get_or_add_rPr() if hasattr(element, "get_or_add_rPr") else element
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    rPr.append(shd)


def style_run(run, ascii_font, cn_font, size=None, bold=None, italic=None, color=None):
    run.font.name = ascii_font
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), ascii_font)
    rFonts.set(qn("w:hAnsi"), ascii_font)
    rFonts.set(qn("w:eastAsia"), cn_font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_inline(par, text, size=10.5, cn_font=BODY_CN, bold_all=False):
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            r = par.add_run(part[2:-2])
            style_run(r, BODY_ASCII, cn_font, size, bold=True)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            r = par.add_run(part[1:-1])
            style_run(r, MONO, MONO, size - 1.5)
            shade(r._element, "F1F1F1")
        elif LINK_RE.match(part):
            m = LINK_RE.match(part)
            label, target = m.group(1), m.group(2)
            if not target or target.startswith("#"):
                r = par.add_run(label)
                style_run(r, BODY_ASCII, cn_font, size, bold=True)
            else:
                r = par.add_run("%s (%s)" % (label, target))
                style_run(r, BODY_ASCII, cn_font, size)
        else:
            r = par.add_run(part)
            style_run(r, BODY_ASCII, cn_font, size, bold=True if bold_all else None)
    return par


def add_field(par, instr):
    r = par.add_run()
    fld = OxmlElement("w:fldChar"); fld.set(qn("w:fldCharType"), "begin")
    r._element.append(fld)
    r2 = par.add_run()
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
    r2._element.append(it)
    r3 = par.add_run()
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
    r3._element.append(fld2)
    for rr in (r, r2, r3):
        style_run(rr, BODY_ASCII, BODY_CN, 9)


def build():
    with open(SRC, encoding="utf-8") as f:
        lines = f.read().split("\n")

    doc = Document()

    # 页面设置
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(PAGE_W), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(MARGIN)
    sec.top_margin = sec.bottom_margin = Cm(2.2)

    # 正文默认字体
    normal = doc.styles["Normal"]
    normal.font.name = BODY_ASCII
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_CN)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.35

    for name, size, cn in (
        ("Heading 1", 16, HEAD_CN),
        ("Heading 2", 14, HEAD_CN),
        ("Heading 3", 12, HEAD_CN),
        ("Heading 4", 11, HEAD_CN),
    ):
        st = doc.styles[name]
        st.font.name = HEAD_ASCII
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), cn)
        st.paragraph_format.space_before = Pt(10)
        st.paragraph_format.space_after = Pt(6)

    # 页脚页码
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(footer, " PAGE ")

    i = 0
    n = len(lines)
    h1_seen = False

    while i < n:
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        # 分隔线
        if re.fullmatch(r"-{3,}", stripped):
            i += 1
            continue

        # 纯 HTML 标签行（Markdown 内跳转锚点，DOCX 中无意义）
        if re.fullmatch(r"(<[^>]+>)+", stripped):
            i += 1
            continue

        # 代码块
        if stripped.startswith("```"):
            i += 1
            buf = []
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            par = doc.add_paragraph()
            pf = par.paragraph_format
            pf.left_indent = Cm(0.15)
            pf.space_before = Pt(4)
            pf.space_after = Pt(8)
            pf.line_spacing = 1.0
            shade(par._element.get_or_add_pPr(), "F5F5F5")
            for k, cl in enumerate(buf):
                if k:
                    par.add_run().add_break()
                r = par.add_run(cl.replace("\t", "    "))
                style_run(r, MONO, MONO, 8.2)
            continue

        # 表格
        if stripped.startswith("|"):
            block = []
            while i < n and lines[i].strip().startswith("|"):
                block.append(lines[i].strip())
                i += 1
            rows = []
            for b in block:
                if re.fullmatch(r"\|[\s:\-|]+\|", b):
                    continue
                cells = [c.strip() for c in b.strip("|").split("|")]
                rows.append(cells)
            if rows:
                ncol = max(len(r) for r in rows)
                for r in rows:
                    r += [""] * (ncol - len(r))
                weights = []
                for c in range(ncol):
                    w = max(len(r[c]) for r in rows)
                    weights.append(max(w, 6) ** 0.85)
                tot = sum(weights)
                widths = [max(1.5, CONTENT_W * w / tot) for w in weights]
                scale = CONTENT_W / sum(widths)
                widths = [w * scale for w in widths]

                table = doc.add_table(rows=0, cols=ncol)
                table.style = "Table Grid"
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.autofit = False
                for ri, rowdata in enumerate(rows):
                    cells = table.add_row().cells
                    for ci, txt in enumerate(rowdata):
                        cell = cells[ci]
                        cell.width = Cm(widths[ci])
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_after = Pt(1)
                        p.paragraph_format.space_before = Pt(1)
                        p.paragraph_format.line_spacing = 1.1
                        add_inline(p, txt, size=8.8, bold_all=(ri == 0))
                        if ri == 0:
                            shade(cell._tc.get_or_add_tcPr(), "E8EEF7")
                    if ri == 0:
                        trPr = table.rows[0]._tr.get_or_add_trPr()
                        th = OxmlElement("w:tblHeader")
                        trPr.append(th)
                doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        # 标题
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            if level == 1 and not h1_seen:
                h1_seen = True
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(14)
                add_inline(p, text, size=22, cn_font=HEAD_CN, bold_all=True)
            else:
                h = doc.add_heading("", level=min(level, 4) if level > 1 else 1)
                add_inline(h, text, size={1: 16, 2: 14, 3: 12, 4: 11}.get(level, 11),
                           cn_font=HEAD_CN, bold_all=True)
            i += 1
            continue

        # 引用
        if stripped.startswith(">"):
            body = stripped.lstrip(">").strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(6)
            add_inline(p, body, size=10, cn_font="楷体")
            i += 1
            continue

        # 列表
        m = re.match(r"^(\s*)[-*]\s+(.*)$", line)
        if m:
            indent = len(m.group(1))
            txt = m.group(2)
            style = "List Bullet" if indent < 2 else "List Bullet 2"
            try:
                p = doc.add_paragraph(style=style)
            except KeyError:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.75 + (0.6 if indent >= 2 else 0))
                txt = "• " + txt
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.3
            add_inline(p, txt)
            i += 1
            continue

        m = re.match(r"^(\s*)(\d+)\.\s+(.*)$", line)
        if m:
            pad = len(m.group(1))
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.9 + (0.6 if pad >= 2 else 0))
            p.paragraph_format.first_line_indent = Cm(-0.55)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.3
            add_inline(p, "%s. %s" % (m.group(2), m.group(3)))
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        add_inline(p, stripped)
        i += 1

    doc.save(OUT)
    print("saved:", OUT)


if __name__ == "__main__":
    build()
